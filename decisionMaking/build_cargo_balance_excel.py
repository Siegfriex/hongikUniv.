"""
실습문제 #3: 화물 적재·선박 밸런스 문제
축구공 엑셀과 동일한 서식 스타일. 답: 총이익 $13,330
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

wb = openpyxl.Workbook()

# ── 스타일 (축구공과 동일) ──
bold = Font(bold=True)
bold_red = Font(bold=True, color="FF0000")
bold_blue = Font(bold=True, color="0000FF")
bold_green = Font(bold=True, color="006600")
title_font = Font(bold=True, size=14, color="FF0000")
section_font = Font(bold=True, size=11, color="333399")
header_font = Font(bold=True, size=12)

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
light_gray = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
light_blue = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
light_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
light_red = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")

thin = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
blue_bd = Border(left=Side("medium", color="0000FF"), right=Side("medium", color="0000FF"),
                 top=Side("medium", color="0000FF"), bottom=Side("medium", color="0000FF"))
red_bd = Border(left=Side("medium", color="FF0000"), right=Side("medium", color="FF0000"),
                top=Side("medium", color="FF0000"), bottom=Side("medium", color="FF0000"))
ctr = Alignment(horizontal="center", vertical="center")
wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
num2 = "0.00"
num0 = "#,##0"
num_comma = "#,##0.00"


def sc(ws, r, c, val, font=None, fill=None, align=None, fmt=None, border=None):
    cell = ws.cell(row=r, column=c, value=val)
    if font: cell.font = font
    if fill: cell.fill = fill
    if align: cell.alignment = align
    if fmt: cell.number_format = fmt
    if border: cell.border = border
    return cell


# ════════════════════════════════════════════════════════════
# Sheet 1: 모형
# ════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "모형"

ws.column_dimensions["A"].width = 22
for c in "BCDE":
    ws.column_dimensions[c].width = 12
ws.column_dimensions["F"].width = 14
ws.column_dimensions["G"].width = 45

# ── 행 1: 제목 ──
sc(ws, 1, 1, "실습 #3  화물 적재·선박 밸런스", font=title_font)

# ── 행 3~7: 화물 데이터 ──
sc(ws, 3, 1, "화물 데이터", font=header_font)

sc(ws, 4, 2, "무게(tons)", font=bold, fill=light_gray, align=ctr)
sc(ws, 4, 3, "부피(cf/ton)", font=bold, fill=light_gray, align=ctr)
sc(ws, 4, 4, "이익($/ton)", font=bold, fill=light_gray, align=ctr)

cargos = [("A", 20, 500, 320), ("B", 16, 700, 400), ("C", 25, 600, 360), ("D", 13, 400, 290)]
for i, (name, wt, vol, profit) in enumerate(cargos):
    sc(ws, 5+i, 1, f"Cargo {name}", font=bold)
    sc(ws, 5+i, 2, wt, align=ctr, border=thin)
    sc(ws, 5+i, 3, vol, align=ctr, border=thin)
    sc(ws, 5+i, 4, profit, align=ctr, border=thin)

sc(ws, 5, 5, "← B5:B8", font=bold_green)
sc(ws, 6, 5, "← C5:C8", font=bold_green)
sc(ws, 7, 5, "← D5:D8", font=bold_green)

# ── 행 10~13: 선박 데이터 ──
sc(ws, 10, 1, "선박 구역", font=header_font)

sc(ws, 11, 2, "최대무게(tons)", font=bold, fill=light_gray, align=ctr)
sc(ws, 11, 3, "최대부피(cf)", font=bold, fill=light_gray, align=ctr)

positions = [("앞(Front)", 12, 7000), ("중간(Mid)", 18, 9000), ("뒤(Rear)", 10, 5000)]
for i, (name, mw, mv) in enumerate(positions):
    sc(ws, 12+i, 1, name, font=bold)
    sc(ws, 12+i, 2, mw, align=ctr, border=thin)
    sc(ws, 12+i, 3, mv, align=ctr, border=thin, fmt=num0)

# ── 행 16~22: 결정변수 매트릭스 ──
sc(ws, 16, 1, "적재 계획", font=header_font)
sc(ws, 16, 7, "【결정변수 4×3 = 12개】", font=section_font)

sc(ws, 17, 2, "앞(F)", font=bold, fill=light_gray, align=ctr)
sc(ws, 17, 3, "중간(M)", font=bold, fill=light_gray, align=ctr)
sc(ws, 17, 4, "뒤(R)", font=bold, fill=light_gray, align=ctr)
sc(ws, 17, 5, "행합", font=bold, fill=light_gray, align=ctr)
sc(ws, 17, 6, "가용량", font=bold, fill=light_gray, align=ctr)

# 행 18~21: 결정변수 (초기값 0)
cargo_names = ["Cargo A", "Cargo B", "Cargo C", "Cargo D"]
cargo_avail = [20, 16, 25, 13]

for i in range(4):
    sc(ws, 18+i, 1, cargo_names[i], font=bold)
    for j in range(3):
        sc(ws, 18+i, 2+j, 0, font=bold_blue, fill=light_blue, align=ctr, border=blue_bd, fmt=num2)
    # 행합
    r = 18 + i
    sc(ws, r, 5, f"=SUM(B{r}:D{r})", align=ctr, fmt=num2, border=thin)
    # 가용량
    sc(ws, r, 6, f"=$B${5+i}", align=ctr, border=thin)

sc(ws, 18, 7, "← ★결정변수 x_ij (Solver 변경 셀)", font=bold_blue)
sc(ws, 19, 7, "← E열=행합(화물별 적재량) ≤ F열(가용량)", font=bold_green)

# ── 행 23~24: 구역별 무게 합 (열합) ──
sc(ws, 23, 1, "구역별 무게합(tons)", font=bold_red)
for j in range(3):
    col = ["B", "C", "D"][j]
    sc(ws, 23, 2+j, f"=SUM({col}18:{col}21)", align=ctr, fmt=num2, border=red_bd, fill=light_red)
sc(ws, 23, 7, "← 열합: 각 구역에 실린 총 무게 (무게 LHS)", font=bold_red)

sc(ws, 24, 2, "<=", align=ctr)
sc(ws, 24, 3, "<=", align=ctr)
sc(ws, 24, 4, "<=", align=ctr)
sc(ws, 24, 7, "← 제약②: 구역 무게합 ≤ 최대무게", font=bold_green)

sc(ws, 25, 1, "최대허용무게(tons)", font=bold)
sc(ws, 25, 2, f"=$B$12", align=ctr, border=thin)
sc(ws, 25, 3, f"=$B$13", align=ctr, border=thin)
sc(ws, 25, 4, f"=$B$14", align=ctr, border=thin)

# ── 행 27~29: 구역별 부피 합 ──
sc(ws, 27, 1, "구역별 부피합(cf)", font=bold_red)
# SUMPRODUCT(부피벡터, 해당열)
sc(ws, 27, 2, "=SUMPRODUCT($C$5:$C$8,B18:B21)", align=ctr, fmt=num0, border=red_bd, fill=light_red)
sc(ws, 27, 3, "=SUMPRODUCT($C$5:$C$8,C18:C21)", align=ctr, fmt=num0, border=red_bd, fill=light_red)
sc(ws, 27, 4, "=SUMPRODUCT($C$5:$C$8,D18:D21)", align=ctr, fmt=num0, border=red_bd, fill=light_red)
sc(ws, 27, 7, "← SUMPRODUCT(부피/ton, 적재톤수) = 실제 부피", font=bold_red)

sc(ws, 28, 2, "<=", align=ctr)
sc(ws, 28, 3, "<=", align=ctr)
sc(ws, 28, 4, "<=", align=ctr)
sc(ws, 28, 7, "← 제약③: 구역 부피합 ≤ 최대부피", font=bold_green)

sc(ws, 29, 1, "최대허용부피(cf)", font=bold)
sc(ws, 29, 2, f"=$C$12", align=ctr, border=thin, fmt=num0)
sc(ws, 29, 3, f"=$C$13", align=ctr, border=thin, fmt=num0)
sc(ws, 29, 4, f"=$C$14", align=ctr, border=thin, fmt=num0)

# ── 행 31~34: 밸런스 제약 ──
sc(ws, 31, 1, "밸런스 제약", font=header_font)
sc(ws, 31, 7, "【핵심! 비율 = 비율 → 교차곱 → 선형 등식】", font=section_font)

sc(ws, 32, 1, "무게/최대무게 비율", font=bold)
sc(ws, 32, 2, "=B23/B25", align=ctr, fmt="0.0000", border=thin)
sc(ws, 32, 3, "=C23/C25", align=ctr, fmt="0.0000", border=thin)
sc(ws, 32, 4, "=D23/D25", align=ctr, fmt="0.0000", border=thin)
sc(ws, 32, 7, "← 앞/12, 중/18, 뒤/10. 이 셋이 같아야 함", font=bold_green)

sc(ws, 33, 1, "밸런스 LHS (앞=중)", font=bold)
sc(ws, 33, 2, "=3*B23-2*C23", align=ctr, fmt=num2, border=thin, fill=light_green)
sc(ws, 33, 3, "=", font=bold_green, align=ctr)
sc(ws, 33, 4, 0, align=ctr, border=thin)
sc(ws, 33, 7, "← 3·Σ앞 - 2·Σ중 = 0  (18·Σ앞 = 12·Σ중 의 약분)", font=bold_green)

sc(ws, 34, 1, "밸런스 LHS (중=뒤)", font=bold)
sc(ws, 34, 2, "=5*C23-9*D23", align=ctr, fmt=num2, border=thin, fill=light_green)
sc(ws, 34, 3, "=", font=bold_green, align=ctr)
sc(ws, 34, 4, 0, align=ctr, border=thin)
sc(ws, 34, 7, "← 5·Σ중 - 9·Σ뒤 = 0  (10·Σ중 = 18·Σ뒤 의 약분)", font=bold_green)

# ── 행 36~37: 목적 셀 ──
sc(ws, 36, 1, "목적함수", font=header_font)

sc(ws, 37, 1, "총 이익 ($)", font=Font(bold=True, size=12))
tc = sc(ws, 37, 2, "=SUMPRODUCT($D$5:$D$8,E18:E21)", font=Font(bold=True, size=12, color="0000FF"),
        fmt=num_comma, align=ctr)
tc.border = Border(left=Side("double", color="0000FF"), right=Side("double", color="0000FF"),
                   top=Side("double", color="0000FF"), bottom=Side("double", color="0000FF"))
sc(ws, 37, 7, "← ★목적셀: Σ(이익/ton × 화물별 적재량) → Max. 답=$13,330", font=bold_blue)

# ── 행 39~: 수식 참조 ──
sc(ws, 39, 1, "수식 참조", font=section_font)
refs = [
    "E18 = SUM(B18:D18)              행합: 화물A 적재량",
    "B23 = SUM(B18:B21)              열합: 앞 구역 총 무게",
    "B27 = SUMPRODUCT($C$5:$C$8, B18:B21)  앞 구역 총 부피",
    "B32 = B23/B25                   비율 (확인용, 제약 아님)",
    "B33 = 3*B23 - 2*C23             밸런스 LHS (= 0이어야)",
    "B34 = 5*C23 - 9*D23             밸런스 LHS (= 0이어야)",
    "B37 = SUMPRODUCT($D$5:$D$8, E18:E21)  총이익 (목적 셀)",
]
for i, ref in enumerate(refs):
    sc(ws, 40+i, 1, ref, font=bold_red)

# ── 행 48~: Solver 설정 ──
sc(ws, 48, 1, "Solver 설정", font=section_font)
solver = [
    ("목적 셀", "B37", "→ Max"),
    ("변경 셀", "B18:D21 (12개)", "x_ij (화물×구역)"),
    ("제약①", "E18:E21 ≤ F18:F21", "화물 가용량"),
    ("제약②", "B23:D23 ≤ B25:D25", "구역 무게 상한"),
    ("제약③", "B27:D27 ≤ B29:D29", "구역 부피 상한"),
    ("제약④", "B33 = D33 (=0)", "밸런스 앞=중"),
    ("제약⑤", "B34 = D34 (=0)", "밸런스 중=뒤"),
    ("비음", "B18:D21 ≥ 0", "Solver 옵션"),
    ("해법", "Simplex LP", "답: $13,330"),
]
for i, (item, setting, note) in enumerate(solver):
    sc(ws, 49+i, 1, item, font=bold)
    sc(ws, 49+i, 2, setting, align=ctr)
    sc(ws, 49+i, 4, note, font=bold_green)


# ════════════════════════════════════════════════════════════
# Sheet 2: 수식해설
# ════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("수식해설")
ws2.column_dimensions["A"].width = 10
ws2.column_dimensions["B"].width = 38
ws2.column_dimensions["C"].width = 38
ws2.column_dimensions["D"].width = 55

r = 1
sc(ws2, r, 1, "수식 해설 — 화물·밸런스 문제", font=Font(bold=True, size=14, color="333399")); r += 2

for j, h in enumerate(["셀", "수식", "계산 내용", "인과"], 1):
    sc(ws2, r, j, h, font=bold, fill=light_gray, align=ctr)
r += 1

fmls = [
    ("B18:D21", "★ 결정변수", "화물i를 구역j에 싣는 톤수 x_ij", "Solver 변경 셀. 4×3=12개"),
    ("E18", "=SUM(B18:D18)", "화물A 적재 총량 (행합)", "화물 가용 제약 LHS. ≤ 20"),
    ("B23", "=SUM(B18:B21)", "앞 구역 총 무게 (열합)", "구역 무게 제약 LHS. ≤ 12"),
    ("B27", "=SUMPRODUCT($C$5:$C$8,B18:B21)", "앞 구역 총 부피", "500·xAF+700·xBF+600·xCF+400·xDF. ≤ 7000"),
    ("B32", "=B23/B25", "비율 (앞무게/앞최대)", "확인용. 최적해에서 B32=C32=D32이면 밸런스 OK"),
    ("B33", "=3*B23-2*C23", "밸런스 LHS: 3Σ앞-2Σ중", "=0이어야. 18Σ앞=12Σ중을 약분한 형태"),
    ("B34", "=5*C23-9*D23", "밸런스 LHS: 5Σ중-9Σ뒤", "=0이어야. 10Σ중=18Σ뒤를 약분한 형태"),
    ("B37", "=SUMPRODUCT($D$5:$D$8,E18:E21)", "총이익 = Σ(이익단가×적재량)", "★ 목적 셀 → Max. 답: $13,330"),
]

for rd in fmls:
    for j, val in enumerate(rd, 1):
        sc(ws2, r, j, val, font=bold_blue if j == 1 else None, align=wrap)
    r += 1

r += 2
sc(ws2, r, 1, "밸런스 제약 도출 과정", font=Font(bold=True, size=12, color="333399")); r += 1

bal = [
    "원래 조건: 앞무게/12 = 중무게/18 = 뒤무게/10",
    "",
    "(예) 앞=6, 중=9, 뒤=5 → 6/12=0.5, 9/18=0.5, 5/10=0.5 OK",
    "",
    "등식 2개로 분리:",
    "  앞/12 = 중/18  →  18·앞 = 12·중  →  3·앞 - 2·중 = 0",
    "  중/18 = 뒤/10  →  10·중 = 18·뒤  →  5·중 - 9·뒤 = 0",
    "",
    "W사 품질 제약과의 비교:",
    "  W사: 가중평균 ≥ 기준 → 분수→선형 '부등식'",
    "  여기: 비율 = 비율 → 교차곱→선형 '등식'",
    "  같은 '비율 선형화' 기법이되, 부등식 vs 등식이 다름",
]

for line in bal:
    sc(ws2, r, 1, line, font=Font(name="Consolas", size=10))
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1


# ════════════════════════════════════════════════════════════
# Sheet 3: 제약조건
# ════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("제약조건")
ws3.column_dimensions["A"].width = 6
ws3.column_dimensions["B"].width = 18
ws3.column_dimensions["C"].width = 35
ws3.column_dimensions["D"].width = 16
ws3.column_dimensions["E"].width = 55

r = 1
sc(ws3, r, 1, "제약조건 해설", font=Font(bold=True, size=14, color="333399")); r += 2

for j, h in enumerate(["#", "제약 이름", "수식", "엑셀 셀", "인과"], 1):
    sc(ws3, r, j, h, font=bold, fill=light_gray, align=ctr)
r += 1

cons = [
    ("①", "화물 가용량", "Σ_j x_ij ≤ W_i (4개)", "E18:E21 ≤ F18:F21",
     "화물 보유톤수 초과 불가. 전량 실을 필요 없음(≤)", light_blue),
    ("②", "구역 무게", "Σ_i x_ij ≤ C_j^w (3개)", "B23:D23 ≤ B25:D25",
     "구역별 적재 무게 한도", light_blue),
    ("③", "구역 부피", "Σ_i v_i·x_ij ≤ C_j^v (3개)", "B27:D27 ≤ B29:D29",
     "부피/ton이 다르므로 SUMPRODUCT로 계산", light_blue),
    ("④", "밸런스(앞=중)", "3·Σ앞 - 2·Σ중 = 0", "B33 = 0",
     "★ 선박 균형. 비율→교차곱→등식", light_green),
    ("⑤", "밸런스(중=뒤)", "5·Σ중 - 9·Σ뒤 = 0", "B34 = 0",
     "★ 독립 등식 2개로 3구역 비율 동일 보장", light_green),
    ("⑥", "비음", "x_ij ≥ 0", "B18:D21 ≥ 0",
     "음수 적재 불가", light_red),
]

for num, name, formula, cell_ref, reason, fl in cons:
    for j, val in enumerate([num, name, formula, cell_ref, reason], 1):
        sc(ws3, r, j, val, align=wrap, fill=fl)
    r += 1

r += 2
sc(ws3, r, 1, "Solver 설정", font=Font(bold=True, size=12, color="333399")); r += 1
for j, h in enumerate(["항목", "설정", "비고"], 1):
    sc(ws3, r, j, h, font=bold, fill=light_gray, align=ctr)
r += 1

for item, setting, note in [
    ("목적 셀", "B37", "→ Max"),
    ("변경 셀", "B18:D21 (12개)", "x_ij"),
    ("제약①", "E18:E21 ≤ F18:F21", "화물 가용"),
    ("제약②", "B23:D23 ≤ B25:D25", "구역 무게"),
    ("제약③", "B27:D27 ≤ B29:D29", "구역 부피"),
    ("제약④", "B33 = D33", "밸런스 앞=중 (=0)"),
    ("제약⑤", "B34 = D34", "밸런스 중=뒤 (=0)"),
    ("비음", "B18:D21 ≥ 0", "Solver 옵션"),
    ("해법", "Simplex LP", "답: $13,330"),
]:
    sc(ws3, r, 1, item, font=bold)
    sc(ws3, r, 2, setting, align=wrap)
    sc(ws3, r, 3, note, align=wrap)
    r += 1


# ════════════════════════════════════════════════════════════
# 저장
# ════════════════════════════════════════════════════════════
out = "/home/sieg/projects-wsl/hongikUniv.-26_1/decisionMaking/실습3_화물적재_밸런스.xlsx"
wb.save(out)
print(f"✓ 저장 완료: {out}")
print(f"  시트: {wb.sheetnames}")
print("\n결정변수 초기값=0. Solver 실행 → 답: 총이익 $13,330")
