"""
행복가구 총괄생산계획 (예제 3.8)
실수 최적해 + 정수 최적해 + 수식해설 + 제약조건 시트
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── 스타일 ──
bold = Font(bold=True)
bold_red = Font(bold=True, color="FF0000")
bold_blue = Font(bold=True, color="0000FF")
bold_green = Font(bold=True, color="006600")
bold_purple = Font(bold=True, color="660099")
title_font = Font(bold=True, size=14, color="FF0000")
section_font = Font(bold=True, size=12)
header_font = Font(bold=True, size=11, color="333399")

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
light_yellow = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
light_gray = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
light_blue = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
light_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
light_red = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")

thin = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
blue_bd = Border(left=Side("medium", color="0000FF"), right=Side("medium", color="0000FF"),
                 top=Side("medium", color="0000FF"), bottom=Side("medium", color="0000FF"))
red_bd = Border(left=Side("medium", color="FF0000"), right=Side("medium", color="FF0000"),
                top=Side("medium", color="FF0000"), bottom=Side("medium", color="FF0000"))
ctr = Alignment(horizontal="center", vertical="center")
wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)


def sc(ws, r, c, val, font=None, fill=None, align=None, fmt=None, border=None):
    cell = ws.cell(row=r, column=c, value=val)
    if font: cell.font = font
    if fill: cell.fill = fill
    if align: cell.alignment = align
    if fmt: cell.number_format = fmt
    if border: cell.border = border
    return cell


def build_model_sheet(ws, sheet_title, hire_vals, fire_vals, ot_vals, prod_vals, is_integer=False):
    """모형 시트를 구축 (실수/정수 공통 레이아웃, 값만 다름)"""

    # 열 너비
    ws.column_dimensions["A"].width = 28
    for c in "BCDEF":
        ws.column_dimensions[c].width = 12
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 50

    num2 = "0.00"
    num0 = "#,##0"
    num_comma = "#,##0.00"
    months = ["1월", "2월", "3월", "4월", "5월"]
    demands = [350, 400, 550, 200, 210]

    # ═══ 행 1: 제목 ═══
    sc(ws, 1, 1, f"행복가구의 총괄계획({sheet_title})", title_font)

    # ═══ 행 3~14: 입력자료 ═══
    sc(ws, 3, 1, "입력자료", section_font)
    sc(ws, 3, 8, "【파라미터 영역】", header_font)

    params = [
        (4,  "초기재고",                        55,  "개",   "I₀ = 55. 재고 Balance Eq.의 시작점"),
        (5,  "초기 작업자수",                    3,   "명",   "W₀ = 3. 인원 Balance Eq.의 시작점"),
        (6,  "정규작업시간/월간/작업자당",       160, "시간", "1인당 월 정규시간. 정규작업시간 = W_t × 160"),
        (7,  "초과작업시간/월간/작업자당",       20,  "시간", "1인당 월 최대 초과. 제약: O_t ≤ W_t × 20"),
        (8,  "고용비용(작업자당)",               170, "만원", "목적계수. 고용비 = 170 × H_t"),
        (9,  "해고비용(작업자당)",               210, "만원", "목적계수. 해고비 = 210 × F_t (고용보다 비쌈!)"),
        (10, "정규임금/월간/작업자당",           150, "만원", "목적계수. 정규임금 = 150 × W_t"),
        (11, "초과작업 수당(시간당)",            2,   "만원", "목적계수. 초과수당 = 2 × O_t"),
        (12, "작업요구시간(옷장당)",             4,   "시간", "생산능력 = 총작업시간 / 4"),
        (13, "재료비(옷장당)",                   3,   "만원", "목적계수. 재료비 = 3 × P_t"),
        (14, "재고유지비(옷장당)",               2,   "만원", "목적계수. 재고비 = 2 × I_t"),
    ]

    for row, name, val, unit, note in params:
        sc(ws, row, 1, name, bold)
        c = sc(ws, row, 2, val, bold, border=thin)
        if row in (4, 5):
            c.fill = light_yellow
        sc(ws, row, 3, unit)
        sc(ws, row, 8, f"← {note}", bold_green)

    # ═══ 행 16~21: 작업자 고용 계획 ═══
    sc(ws, 16, 1, "작업자 고용 계획", section_font)
    sc(ws, 16, 3, "고용/해고: 월초에 발생한다 가정", bold_red)
    sc(ws, 16, 8, "【인원 Balance Equation 영역】", header_font)

    # 행 17: 월 헤더
    for i, m in enumerate(months):
        sc(ws, 17, 2 + i, m, bold, light_gray, ctr)

    # 행 18: 이월되는 작업자수
    sc(ws, 18, 1, "이월되는 작업자수", bold)
    sc(ws, 18, 2, "=$B$5", align=ctr, border=thin, fmt=num2)  # B18 = 초기 작업자수
    for i in range(1, 5):
        cl = get_column_letter(1 + i)
        sc(ws, 18, 2 + i, f"={cl}21", align=ctr, border=thin, fmt=num2)  # C18=B21, D18=C21...
    sc(ws, 18, 8, "← B18=$B$5(초기인원). C18=B21(전월 근무가능=이번달 이월)", bold_green)

    # 행 19: 고용 작업자수 — ★ 결정변수
    sc(ws, 19, 1, "고용 작업자수", bold)
    for i, h in enumerate(hire_vals):
        c = sc(ws, 19, 2 + i, h, bold_blue, light_blue, ctr, fmt=num2)
        c.border = blue_bd
    sc(ws, 19, 8, "← ★ 결정변수 H_t (Solver 변경 셀)", bold_blue)

    # 행 20: 해고 작업자수 — ★ 결정변수
    sc(ws, 20, 1, "해고 작업자수", bold)
    for i, f in enumerate(fire_vals):
        c = sc(ws, 20, 2 + i, f, bold_blue, light_blue, ctr, fmt=num2)
        c.border = blue_bd
    sc(ws, 20, 8, "← ★ 결정변수 F_t (Solver 변경 셀)", bold_blue)

    # 행 21: 근무가능한 작업자수 — Balance Eq.
    sc(ws, 21, 1, "근무가능한 작업자수", bold_red)
    for i in range(5):
        cl = get_column_letter(2 + i)
        formula = f"={cl}18+{cl}19-{cl}20"
        c = sc(ws, 21, 2 + i, formula, align=ctr, fmt=num2)
        c.border = red_bd
        c.fill = light_red
    sc(ws, 21, 7, "W_t", bold_purple)
    sc(ws, 21, 8, "← ★ 인원 Balance: W_t = W_{t-1} + H_t - F_t", bold_red)

    # ═══ 행 23~28: 작업시간 ═══
    # 행 23: 가능한 정규작업시간
    sc(ws, 23, 1, "가능한 정규작업시간", bold)
    for i in range(5):
        cl = get_column_letter(2 + i)
        sc(ws, 23, 2 + i, f"={cl}21*$B$6", align=ctr, fmt=num2, border=thin)
    sc(ws, 23, 8, "← W_t × 160. 인원이 정하면 자동 결정", bold_green)

    # 행 24: 초과작업시간 — ★ 결정변수
    sc(ws, 24, 1, "초과작업시간", bold)
    for i, o in enumerate(ot_vals):
        c = sc(ws, 24, 2 + i, o, bold_blue, light_blue, ctr, fmt=num2)
        c.border = blue_bd
    sc(ws, 24, 8, "← ★ 결정변수 O_t (Solver 변경 셀)", bold_blue)

    # 행 25: <= 표시
    for i in range(5):
        sc(ws, 25, 2 + i, "<=", align=ctr)
    sc(ws, 25, 8, "← 제약②: 초과작업 ≤ 최대 초과작업 가능시간", bold_green)

    # 행 26: 최대 초과작업 가능시간
    sc(ws, 26, 1, "최대 초과작업 가능시간", bold)
    for i in range(5):
        cl = get_column_letter(2 + i)
        sc(ws, 26, 2 + i, f"={cl}21*$B$7", align=ctr, fmt=num2, border=thin)
    sc(ws, 26, 8, "← W_t × 20. 제약②의 RHS", bold_green)

    # 행 28: 생산에 투자된 총작업시간
    sc(ws, 28, 1, "생산에 투자된 총작업시간", bold)
    for i in range(5):
        cl = get_column_letter(2 + i)
        sc(ws, 28, 2 + i, f"={cl}23+{cl}24", align=ctr, fmt=num2, border=thin)
    sc(ws, 28, 8, "← 정규시간 + 초과시간. 생산능력 계산의 입력", bold_green)

    # ═══ 행 30~39: 생산계획 ═══
    sc(ws, 30, 1, "생산계획", section_font)
    sc(ws, 30, 8, "【재고 Balance Equation 영역】", header_font)

    for i, m in enumerate(months):
        sc(ws, 31, 2 + i, m, bold, light_gray, ctr)

    # 행 32: 옷장 생산량 — ★ 결정변수
    sc(ws, 32, 1, "옷장 생산량", bold)
    for i, p in enumerate(prod_vals):
        c = sc(ws, 32, 2 + i, p, bold_blue, light_blue, ctr, fmt=num2)
        c.border = blue_bd
    sc(ws, 32, 8, "← ★ 결정변수 P_t (Solver 변경 셀)", bold_blue)

    # 행 33: <= 표시
    for i in range(5):
        sc(ws, 33, 2 + i, "<=", align=ctr)
    sc(ws, 33, 8, "← 제약③: 생산량 ≤ 생산능력", bold_green)

    # 행 34: 생산능력
    sc(ws, 34, 1, "생산능력", bold)
    for i in range(5):
        cl = get_column_letter(2 + i)
        sc(ws, 34, 2 + i, f"={cl}28/$B$12", align=ctr, fmt=num2, border=thin)
    sc(ws, 34, 8, "← 총작업시간 / 4(시간/개). 제약③의 RHS", bold_green)

    # 행 36: 수요
    sc(ws, 36, 1, "수요", bold)
    for i, d in enumerate(demands):
        sc(ws, 36, 2 + i, d, bold, light_gray, ctr, border=thin)
    sc(ws, 36, 7, "d_t", bold_purple)
    sc(ws, 36, 8, "← 상수. 3월(550) > 정규 최대(400) → 초과 or 재고 필요", bold_green)

    # 행 37: 월말재고 — Balance Eq.
    sc(ws, 37, 1, "월말재고", bold_red)
    # B37 = 초기재고 + 생산 - 수요
    sc(ws, 37, 2, "=$B$4+B32-B36", align=ctr, fmt=num2)
    ws["B37"].border = red_bd
    ws["B37"].fill = light_red
    for i in range(1, 5):
        cl = get_column_letter(2 + i)
        pcl = get_column_letter(1 + i)
        c = sc(ws, 37, 2 + i, f"={pcl}37+{cl}32-{cl}36", align=ctr, fmt=num2)
        c.border = red_bd
        c.fill = light_red
    sc(ws, 37, 7, "I_t", bold_purple)
    sc(ws, 37, 8, "← ★ 재고 Balance: I_t = I_{t-1} + P_t - d_t", bold_red)

    # 행 38: >= 표시
    for i in range(5):
        sc(ws, 38, 2 + i, ">=", align=ctr)
    sc(ws, 38, 8, "← 제약⑤: 백오더 불허 (월말재고 ≥ 0)", bold_green)

    # 행 39: 수요만족을 위한 조건
    sc(ws, 39, 1, "수요만족을 위한 조건", bold)
    for i in range(5):
        sc(ws, 39, 2 + i, 0, align=ctr, border=thin)
    sc(ws, 39, 8, "← 하한 = 0. I_t ≥ 0", bold_green)

    # ═══ 행 41~49: 발생비용 ═══
    sc(ws, 41, 1, "발생비용(만원)", section_font)
    sc(ws, 41, 8, "【목적함수 영역 — 6가지 비용】", header_font)

    for i, m in enumerate(months):
        sc(ws, 42, 2 + i, m, bold, light_gray, ctr)
    sc(ws, 42, 7, "합계", bold, light_gray, ctr)

    cost_rows = [
        (43, "고용비용",     "19", "$B$8",  "170 × H_t"),
        (44, "해고비용",     "20", "$B$9",  "210 × F_t"),
        (45, "정규임금",     "21", "$B$10", "150 × W_t"),
        (46, "초과작업수당", "24", "$B$11", "2 × O_t"),
        (47, "재료비",       "32", "$B$13", "3 × P_t"),
        (48, "재고유지비",   "37", "$B$14", "2 × I_t"),
    ]

    for row, name, var_row, param_ref, desc in cost_rows:
        sc(ws, row, 1, name, bold)
        for i in range(5):
            cl = get_column_letter(2 + i)
            sc(ws, row, 2 + i, f"={cl}{var_row}*{param_ref}", fmt=num_comma, align=ctr, border=thin)
        # G열: 합계
        sc(ws, row, 7, f"=SUM(B{row}:F{row})", fmt=num_comma, align=ctr, border=thin, font=bold)
        sc(ws, row, 8, f"← {desc}", bold_green)

    # 행 49: 월별 합계
    sc(ws, 49, 1, "합계", bold)
    for i in range(5):
        cl = get_column_letter(2 + i)
        sc(ws, 49, 2 + i, f"=SUM({cl}43:{cl}48)", fmt=num_comma, align=ctr, border=thin, font=bold)

    # G49: 전체비용 (목적 셀)
    tc = sc(ws, 49, 7, "=SUM(G43:G48)", Font(bold=True, size=12), fmt=num_comma, align=ctr)
    tc.border = Border(
        left=Side("double", color="0000FF"), right=Side("double", color="0000FF"),
        top=Side("double", color="0000FF"), bottom=Side("double", color="0000FF"))
    sc(ws, 49, 8, "← ★ 목적 셀: min Z = Σ(고용+해고+정규+초과+재료+재고)", bold_blue)

    # ═══ 행 51~: 수식 주석 ═══
    sc(ws, 51, 1, "수식 참조", section_font)
    notes = [
        "B21 = B18 + B19 - B20  (인원 Balance: 이월 + 고용 - 해고)",
        "C18 = B21  (전월 근무가능 → 이번달 이월)",
        "B23 = B21 * $B$6  (정규시간 = 인원 × 160)",
        "B26 = B21 * $B$7  (최대초과 = 인원 × 20)",
        "B28 = B23 + B24  (총작업시간 = 정규 + 초과)",
        "B34 = B28 / $B$12  (생산능력 = 총시간 / 4)",
        "B37 = $B$4 + B32 - B36  (1월 재고: 초기재고 + 생산 - 수요)",
        "C37 = B37 + C32 - C36  (2월~ 재고: 전월재고 + 생산 - 수요)",
    ]
    for j, note in enumerate(notes):
        sc(ws, 52 + j, 1, note, bold_red)
        ws.merge_cells(start_row=52 + j, start_column=1, end_row=52 + j, end_column=7)


# ════════════════════════════════════════════════════════════
# Sheet 1: 실수 최적해
# ════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "실수최적해"

# 실수 최적해 값 (이미지 1~3에서 읽은 값)
build_model_sheet(ws1, "실수 최적해",
    hire_vals=[6.96, 0, 0, 0, 0],
    fire_vals=[0, 0, 0, 4.84, 0],
    ot_vals=[0, 0, 199.20, 0, 0],
    prod_vals=[398.40, 398.40, 448.20, 205.00, 205.00],
    is_integer=False,
)

# 실수 최적해 추가 주석
sc(ws1, 61, 1, "실수 최적해 결과", Font(bold=True, size=12, color="FF0000"))
sc(ws1, 62, 1, "목표셀(총비용) ≈ 14,001 (만원)", bold_red)
sc(ws1, 63, 1, "→ 다음 단계: 정수 조건 추가 (고용·해고를 정수로)", bold_red)
sc(ws1, 64, 1, "  (1) 정수 제약: B19:F19, B20:F20 = integer", bold)
sc(ws1, 65, 1, "  (2) 옵션: 정수 최적화 비율(허용한도) = 0%", bold)


# ════════════════════════════════════════════════════════════
# Sheet 2: 정수 최적해
# ════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("정수최적해")

# 정수 최적해 값 (이미지 4~5에서 읽은 값)
build_model_sheet(ws2, "정수 최적해",
    hire_vals=[7, 0, 0, 0, 0],
    fire_vals=[0, 0, 0, 5, 0],
    ot_vals=[0, 0, 180, 0, 40],
    prod_vals=[400, 400, 445, 200, 210],
    is_integer=True,
)

# 정수 최적해 추가 주석
sc(ws2, 61, 1, "정수 최적해 결과", Font(bold=True, size=12, color="FF0000"))
sc(ws2, 62, 1, "목표셀(총비용) = 14,065 (만원) ≈ 1억 4,065만원", bold_red)
sc(ws2, 63, 1, "실수해(14,001) → 정수해(14,065): 차이 64만원 = 정수 제약의 비용", bold)


# ════════════════════════════════════════════════════════════
# Sheet 3: 수식해설
# ════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("수식해설")
ws3.column_dimensions["A"].width = 8
ws3.column_dimensions["B"].width = 20
ws3.column_dimensions["C"].width = 35
ws3.column_dimensions["D"].width = 55

r = 1
sc(ws3, r, 1, "수식 해설 — 각 셀이 무엇을 계산하고, 왜 필요한가", Font(bold=True, size=14, color="333399")); r += 2

for j, h in enumerate(["셀", "수식", "무엇을 계산하는가", "왜 필요한가 (인과)"], 1):
    sc(ws3, r, j, h, bold, light_gray, ctr)
r += 1

formulas = [
    ("B18", "=$B$5", "1월초 이월 인원 = 초기 작업자 3명", "인원 Balance의 시작점 (y₀에 해당)"),
    ("C18", "=B21", "2월 이월 = 1월 근무가능 인원", "인원 체인: 전월 말 인원 → 이번달 초 이월"),
    ("B19:F19", "(결정변수)", "월별 고용 H_t (명)", "★ Solver 변경 셀. 인원을 늘리는 레버"),
    ("B20:F20", "(결정변수)", "월별 해고 F_t (명)", "★ Solver 변경 셀. 인원을 줄이는 레버"),
    ("B21", "=B18+B19-B20", "근무가능 W_t = 이월+고용-해고", "★ 인원 Balance Eq. 핵심. W₀=3에서 시작"),
    ("B23", "=B21*$B$6", "정규작업시간 = W_t × 160", "인원이 결정되면 자동. 총작업시간의 기본 부분"),
    ("B24:F24", "(결정변수)", "초과작업시간 O_t", "★ 정규시간으로 부족할 때 추가 생산력 확보"),
    ("B26", "=B21*$B$7", "최대초과 = W_t × 20", "제약②의 RHS. 1인 20h 한도 → 인원에 비례"),
    ("B28", "=B23+B24", "총작업시간 = 정규 + 초과", "생산능력 계산의 입력"),
    ("B34", "=B28/$B$12", "생산능력 = 총시간 / 4", "제약③의 RHS. 축구공(고정30)과 달리 인원에 따라 변동"),
    ("B32:F32", "(결정변수)", "생산량 P_t (개)", "★ 능력 이하로만 가능. 수요보다 많이 만들어 재고 비축 가능"),
    ("B37", "=$B$4+B32-B36", "1월 재고 = 초기(55)+생산-수요", "★ 재고 Balance. 축구공과 동일 구조"),
    ("C37", "=B37+C32-C36", "2월~ 재고 = 전월재고+생산-수요", "재고 체인. I₀=55에서 시작"),
    ("B43", "=B19*$B$8", "고용비 = H_t × 170", "목적함수 6항 중 ①"),
    ("B44", "=B20*$B$9", "해고비 = F_t × 210", "목적함수 ② (고용보다 비쌈!)"),
    ("B45", "=B21*$B$10", "정규임금 = W_t × 150", "목적함수 ③ (인원 비례)"),
    ("B46", "=B24*$B$11", "초과수당 = O_t × 2", "목적함수 ④"),
    ("B47", "=B32*$B$13", "재료비 = P_t × 3", "목적함수 ⑤"),
    ("B48", "=B37*$B$14", "재고비 = I_t × 2", "목적함수 ⑥"),
    ("G49", "=SUM(G43:G48)", "5개월 총비용 Z", "★ 목적 셀 → Solver Min"),
]

for rd in formulas:
    for j, val in enumerate(rd, 1):
        f = bold_blue if j == 1 else None
        sc(ws3, r, j, val, f, align=wrap)
    r += 1

r += 2
sc(ws3, r, 1, "수식 연쇄 흐름도 (2중 Balance)", Font(bold=True, size=12, color="333399")); r += 1

flow = [
    "H_t, F_t (고용·해고) ──→ 인원 Balance ──→ W_t(근무가능인원)",
    "                                            │",
    "                              ┌──── × 160 ──┤── 정규시간(행23)",
    "                              │              │",
    "                              │    ┌── × 20 ─┘── 초과상한(행26, 제약②의 RHS)",
    "                              │    │",
    "                              │    └── O_t(초과, 결정변수) ≤ 초과상한",
    "                              │         │",
    "                              └─ 정규 + O_t ──→ 총시간(행28) ──→ ÷4 ──→ 생산능력(행34, 제약③ RHS)",
    "                                                                           │",
    "P_t(생산, 결정변수) ≤ 생산능력 ─────────────────────────────────────────────┘",
    "     │",
    "     └──→ 재고 Balance ──→ I_t(월말재고) ≥ 0 (제약⑤)",
    "",
    "비용: 170H + 210F + 150W + 2O + 3P + 2I ──→ 총비용 Z(G49) → Min",
]
for line in flow:
    sc(ws3, r, 1, line, Font(name="Consolas", size=9))
    ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1


# ════════════════════════════════════════════════════════════
# Sheet 4: 제약조건
# ════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("제약조건")
ws4.column_dimensions["A"].width = 6
ws4.column_dimensions["B"].width = 20
ws4.column_dimensions["C"].width = 35
ws4.column_dimensions["D"].width = 18
ws4.column_dimensions["E"].width = 55

r = 1
sc(ws4, r, 1, "제약조건 해설", Font(bold=True, size=14, color="333399")); r += 2

for j, h in enumerate(["#", "제약 이름", "수식 (대수)", "엑셀 셀", "인과"], 1):
    sc(ws4, r, j, h, bold, light_gray, ctr)
r += 1

constraints = [
    ("①", "인원 수지", "W_t = W_{t-1}+H_t-F_t (W₀=3)", "행21 수식",
     "엑셀 수식으로 구현. Solver 제약 아님. 고용·해고가 인원으로 이어지는 체인"),
    ("②", "초과작업 상한", "O_t ≤ 20 × W_t", "행24 ≤ 행26",
     "1인당 월 20시간 한도. 인원이 줄면 초과작업 가능량도 줄어듦"),
    ("③", "생산량 상한", "P_t ≤ (160W_t+O_t)/4", "행32 ≤ 행34",
     "총작업시간/4 = 능력. 인원·초과에 따라 변동 (축구공의 고정30과 다름)"),
    ("④", "재고 수지", "I_t = I_{t-1}+P_t-d_t (I₀=55)", "행37 수식",
     "엑셀 수식으로 구현. 축구공(§6)과 동일 구조"),
    ("⑤", "백오더 불허", "I_t ≥ 0", "행37 ≥ 행39",
     "주문미처리 허용 안 함. 당월 수요를 반드시 충족"),
    ("⑥", "비음 조건", "H_t, F_t, O_t, P_t ≥ 0", "변경셀 ≥ 0",
     "음수 고용·해고·초과·생산은 물리적 불가"),
]

fills = [light_red, light_blue, light_blue, light_red, light_green, light_green]
for i, (num, name, formula, cell_ref, reason) in enumerate(constraints):
    for j, val in enumerate([num, name, formula, cell_ref, reason], 1):
        sc(ws4, r, j, val, align=wrap, fill=fills[i])
    r += 1

r += 2
sc(ws4, r, 1, "Solver 설정 요약", Font(bold=True, size=12, color="333399")); r += 1

for j, h in enumerate(["항목", "설정", "비고"], 1):
    sc(ws4, r, j, h, bold, light_gray, ctr)
r += 1

solver = [
    ("목적 셀", "G49 = SUM(G43:G48)", "→ Min (5개월 총비용)"),
    ("변경 셀", "B19:F19, B20:F20, B24:F24, B32:F32", "H,F,O,P × 5 = 20개"),
    ("제약②", "B24:F24 ≤ B26:F26", "초과작업 ≤ 인원×20"),
    ("제약③", "B32:F32 ≤ B34:F34", "생산량 ≤ 생산능력"),
    ("제약⑤", "B37:F37 ≥ 0", "백오더 불허"),
    ("제약⑥", "변경셀 ≥ 0", "비음 (또는 Solver 옵션 체크)"),
    ("해법", "Simplex LP", "1차: 실수 → 14,001"),
    ("", "", ""),
    ("정수화", "B19:F19, B20:F20 = integer", "고용·해고는 정수(사람 수)"),
    ("옵션", "정수 최적화 비율 = 0%", "정확한 정수 최적해 보장"),
    ("결과", "정수 최적해 = 14,065", "실수해 대비 +64만원"),
]

for item, setting, note in solver:
    f = bold_red if item in ("정수화", "옵션") else bold
    sc(ws4, r, 1, item, f)
    sc(ws4, r, 2, setting, align=wrap)
    sc(ws4, r, 3, note, align=wrap)
    r += 1

r += 2
sc(ws4, r, 1, "염두에 둘 것 — 예외 처리·주의사항", Font(bold=True, size=12, color="333399")); r += 1

cautions = [
    "1. 같은 달에 H_t>0, F_t>0 동시 발생? → LP가 자동으로 한쪽만 양수로 만듦 (비용 낭비 방지)",
    "2. 3월 수요(550) 대응: 정규시간만으로 10명×160h/4=400개. 부족 150개 → 초과작업+재고로 보충",
    "   정수해: 1~2월에 미리 만들어 재고 105+105 쌓고, 3월 초과 180h로 445개 생산 → 105+445-550=0",
    "3. 해고비(210) > 고용비(170) → LP는 인원을 안정적으로 유지하려 함",
    "4. 4월 이후 수요 급감(200,210) → 4월에 5명 해고. 해고비 비싸지만 이후 정규임금 절감이 더 큼",
    "5. 정수화 시 주의: 실수해에서 H₁=6.96 → 정수해 7. F₄=4.84 → 5. 이로 인해 Z가 64만원↑",
    "6. Balance Eq.(행21, 행37)는 수식 셀 → Solver 제약에 넣지 않음!",
    "7. 단위: 모든 비용 만원. Z=14,065 → 약 1억 4천만원",
]

for line in cautions:
    sc(ws4, r, 1, line, Font(size=10))
    ws4.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1

r += 2
sc(ws4, r, 1, "7개 문제 전체 대조", Font(bold=True, size=12, color="333399")); r += 1

for j, h in enumerate(["문제", "변수 수", "Balance", "비용 항목", "정수 이슈"], 1):
    sc(ws4, r, j, h, bold, light_gray, ctr)
r += 1

comparison = [
    ("Decora(절단)", "26", "없음", "1(봉 수)", "패턴횟수(정수)"),
    ("Big M(운송)", "6", "없음", "1(운송비)", "없음"),
    ("Sellmore(할당)", "16", "없음", "1(인건비)", "0-1(자동 정수)"),
    ("W사(혼합)", "6", "없음", "1(이익)", "없음"),
    ("축구공(동적)", "6", "1(재고)", "2(생산+재고)", "없음"),
    ("행복가구(총괄)", "20", "2(인원+재고)", "6(고용+해고+임금+초과+재료+재고)", "고용·해고(정수)"),
]

for prob, nvar, bal, cost, intg in comparison:
    sc(ws4, r, 1, prob, bold)
    sc(ws4, r, 2, nvar, align=ctr)
    sc(ws4, r, 3, bal)
    sc(ws4, r, 4, cost)
    sc(ws4, r, 5, intg)
    r += 1


# ════════════════════════════════════════════════════════════
# 저장
# ════════════════════════════════════════════════════════════
out = "/home/sieg/projects-wsl/hongikUniv.-26_1/decisionMaking/예제3-8_행복가구_총괄생산계획.xlsx"
wb.save(out)
print(f"✓ 저장 완료: {out}")
print(f"  시트: {wb.sheetnames}")

# ── 검증 (정수 최적해) ──
print("\n── 검증: 정수 최적해 ──")
H = [7, 0, 0, 0, 0]
F = [0, 0, 0, 5, 0]
O = [0, 0, 180, 0, 40]
P = [400, 400, 445, 200, 210]
D = [350, 400, 550, 200, 210]

W = [0]*5
I = [0]*5
W_prev, I_prev = 3, 55

print(f"{'월':>3} {'W':>4} {'H':>3} {'F':>3} {'정규h':>6} {'OT':>4} {'총h':>5} {'능력':>5} {'P':>4} {'I':>4} | {'고용':>5} {'해고':>5} {'정규':>5} {'초과':>4} {'재료':>5} {'재고':>4} {'합':>6}")
print("-" * 110)

total = 0
for t in range(5):
    W[t] = W_prev + H[t] - F[t]
    I[t] = I_prev + P[t] - D[t]
    reg_h = W[t] * 160
    total_h = reg_h + O[t]
    cap = total_h / 4

    c_hire = H[t] * 170
    c_fire = F[t] * 210
    c_reg = W[t] * 150
    c_ot = O[t] * 2
    c_mat = P[t] * 3
    c_inv = I[t] * 2
    c_sum = c_hire + c_fire + c_reg + c_ot + c_mat + c_inv
    total += c_sum

    print(f"{t+1:>3}월 {W[t]:>4} {H[t]:>3} {F[t]:>3} {reg_h:>6} {O[t]:>4} {total_h:>5} {cap:>5.0f} {P[t]:>4} {I[t]:>4} | {c_hire:>5} {c_fire:>5} {c_reg:>5} {c_ot:>4} {c_mat:>5} {c_inv:>4} {c_sum:>6}")
    W_prev = W[t]
    I_prev = I[t]

print("-" * 110)
print(f"{'총비용':>85} {total:>6}")
